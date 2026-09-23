import json
import os
import openpyxl

ABSENT_KEYWORDS = [
    "NO ASISTE", "NO ASISTIO", "N.A.", "N/A", "RETIRADO", "RETIRADA", 
    "TRASLADADO", "TRASLADADA", "SIN EVALUAR", "S/E", "S.E.", "N.E.", 
    "FALTO", "FALTA", "RETIRO", "EXP. REC", "EXONERADO", "EXONERADA",
    "NO EVALUADO", "NO EVALUADA"
]

def import_bimester_excel(excel_file_path, bimester_name, json_path="evaluaciones_2026.json"):
    """
    Parses an Excel file containing student grade letters (AD, A, B, C) per section and grade,
    updates evaluations_2026.json with the new bimester data.
    Omits absent or non-evaluated students from level C and evaluation statistics.
    """
    if not os.path.exists(excel_file_path):
        return False, f"Archivo {excel_file_path} no encontrado."

    if not os.path.exists(json_path):
        return False, f"Archivo de datos {json_path} no encontrado."

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    wb = openpyxl.load_workbook(excel_file_path, data_only=True)
    sheet = wb.active

    grade_counts = {
        "1° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0},
        "2° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0},
        "3° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0},
        "4° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0},
        "5° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0}
    }

    for r in range(2, sheet.max_row + 1):
        row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
        if not row_vals or all(v is None for v in row_vals):
            continue

        row_text = " ".join(str(v).strip().upper() for v in row_vals if v is not None)
        
        # Omit absent or non-evaluated students completely from C and evaluados statistics
        if any(kw in row_text for kw in ABSENT_KEYWORDS):
            continue

        grade_str = None
        nota_str = None

        for val in row_vals:
            if val is None:
                continue
            v_str = str(val).strip().upper()
            if "GRADO" in v_str or v_str in ["1", "2", "3", "4", "5"]:
                if "1" in v_str: grade_str = "1° Grado"
                elif "2" in v_str: grade_str = "2° Grado"
                elif "3" in v_str: grade_str = "3° Grado"
                elif "4" in v_str: grade_str = "4° Grado"
                elif "5" in v_str: grade_str = "5° Grado"
            
            if v_str in ["AD", "A", "B", "C"]:
                nota_str = v_str
            elif v_str in ["AD/A", "LOGRO"]:
                nota_str = "A"

        if grade_str and nota_str:
            grade_counts[grade_str]["evaluados"] += 1
            if nota_str in ["AD", "A"]:
                grade_counts[grade_str]["AD_A"] += 1
            elif nota_str == "B":
                grade_counts[grade_str]["B"] += 1
            elif nota_str == "C":
                grade_counts[grade_str]["C"] += 1

    tot_eval = 0
    tot_ada = 0
    tot_b = 0
    tot_c = 0

    for g, counts in grade_counts.items():
        e = counts["evaluados"]
        if e > 0:
            counts["pct_AD_A"] = round((counts["AD_A"] / e) * 100, 2)
            counts["pct_B"] = round((counts["B"] / e) * 100, 2)
            counts["pct_C"] = round((counts["C"] / e) * 100, 2)
        else:
            counts["pct_AD_A"] = 0.0
            counts["pct_B"] = 0.0
            counts["pct_C"] = 0.0
            
        tot_eval += e
        tot_ada += counts["AD_A"]
        tot_b += counts["B"]
        tot_c += counts["C"]

        if g not in data["grados"]:
            data["grados"][g] = {}
        data["grados"][g][bimester_name] = counts

    if tot_eval > 0:
        promedio_stats = {
            "evaluados": tot_eval,
            "AD_A": tot_ada,
            "B": tot_b,
            "C": tot_c,
            "pct_AD_A": round((tot_ada / tot_eval) * 100, 2),
            "pct_B": round((tot_b / tot_eval) * 100, 2),
            "pct_C": round((tot_c / tot_eval) * 100, 2)
        }
        data["grados"]["Promedio General"][bimester_name] = promedio_stats

    if bimester_name not in data["meta"]["periodos_disponibles"]:
        data["meta"]["periodos_disponibles"].append(bimester_name)

    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    return True, f"Bimestre '{bimester_name}' importado exitosamente con {tot_eval} estudiantes evaluados."


def import_section_excel(filepath, grado, seccion):
    """
    Parses a single Excel file for one section.
    Flexible parser: scans ALL columns in every row looking for grade letters (AD, A, B, C).
    Omit non-attending / absent / retired / non-evaluated students from 'C' and evaluation statistics.
    """
    try:
        wb = openpyxl.load_workbook(filepath, data_only=True)
        sheet = wb.active
        
        ad_a_count = 0
        b_count = 0
        c_count = 0
        total = 0
        matriculados = 0
        
        # Scan all rows in the sheet
        for r in range(1, sheet.max_row + 1):
            row_vals = [sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)]
            if not row_vals or all(v is None for v in row_vals):
                continue
            
            row_text = " ".join(str(v).strip().upper() for v in row_vals if v is not None)
            
            # Skip pure header rows
            header_keywords = ["ESTUDIANTE", "NOMBRE", "APELLIDO", "N°", "COMPETENCIA", 
                               "CALIFICACION", "NIVEL", "GRADO", "SECCION", "AREA"]
            if any(kw in row_text for kw in header_keywords):
                continue
            
            # Check if this row represents a student
            is_student_row = any(
                isinstance(v, (int, float)) or (isinstance(v, str) and len(v.strip()) > 2)
                for v in row_vals if v is not None
            )
            
            if not is_student_row:
                continue

            # Check if student is explicitly marked as absent / retired / not evaluated
            if any(kw in row_text for kw in ABSENT_KEYWORDS):
                matriculados += 1
                # OMITTED from evaluados (total) and OMITTED from C/AD/B
                continue

            # Scan cells right-to-left to find the grade letter
            grade_found = None
            for c_idx in range(len(row_vals) - 1, -1, -1):
                cell_val = row_vals[c_idx]
                if cell_val is None:
                    continue
                v = str(cell_val).strip().upper()
                if v in ["AD", "A", "B", "C"]:
                    grade_found = v
                    break
                elif v in ["AD/A", "LOGRO"]:
                    grade_found = "A"
                    break
            
            if grade_found:
                matriculados += 1
                total += 1
                if grade_found in ["AD", "A"]:
                    ad_a_count += 1
                elif grade_found == "B":
                    b_count += 1
                elif grade_found == "C":
                    c_count += 1
            else:
                # Student row exists (e.g. student name) but has no valid grade letter
                # Count towards matriculados, but OMIT from evaluados and C level
                matriculados += 1
                    
        pct_ad_a = round((ad_a_count / total) * 100, 2) if total > 0 else 0.0
        pct_b = round((b_count / total) * 100, 2) if total > 0 else 0.0
        pct_c = round((c_count / total) * 100, 2) if total > 0 else 0.0
        
        return {
            "seccion": seccion,
            "grado": grado,
            "matriculados": matriculados,
            "evaluados": total,
            "AD_A": ad_a_count,
            "B": b_count,
            "C": c_count,
            "pct_AD_A": pct_ad_a,
            "pct_B": pct_b,
            "pct_C": pct_c
        }
    except Exception as e:
        print(f"Error leyendo {filepath} para {seccion}: {e}")
        return None


def consolidate_bimester(sections_data, bimester_name, json_path='evaluaciones_2026.json'):
    """
    Takes a list of section results and consolidates them into grade-level
    and institution-level statistics, saving to the JSON database.
    Supports partial uploads (not all 10 sections required).
    """
    if not os.path.exists(json_path):
        return False, f"Archivo de datos {json_path} no encontrado.", {}

    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
        
    if "grados" not in data:
        data["grados"] = {}
    if "secciones_bimestres" not in data:
        data["secciones_bimestres"] = {}
    if "meta" not in data:
        data["meta"] = {"periodos_disponibles": [], "periodos_totales": []}
        
    # Initialize the bimester tracking in secciones_bimestres
    if bimester_name not in data["secciones_bimestres"]:
        data["secciones_bimestres"][bimester_name] = {}
    
    grade_counts = {
        "1° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0, "matriculados": 0},
        "2° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0, "matriculados": 0},
        "3° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0, "matriculados": 0},
        "4° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0, "matriculados": 0},
        "5° Grado": {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0, "matriculados": 0}
    }
    
    total_evaluados = 0
    total_sections = 0
    alerts = []
    processed_sections = []
    
    for sec_data in sections_data:
        if not sec_data:
            continue
            
        grado = sec_data["grado"]
        seccion = sec_data["seccion"]
        
        # Store in secciones_bimestres
        data["secciones_bimestres"][bimester_name][seccion] = {
            "matriculados": sec_data.get("matriculados", sec_data["evaluados"]),
            "evaluados": sec_data["evaluados"],
            "AD_A": sec_data["AD_A"],
            "B": sec_data["B"],
            "C": sec_data["C"],
            "pct_AD_A": sec_data["pct_AD_A"],
            "pct_B": sec_data["pct_B"],
            "pct_C": sec_data["pct_C"],
            "grado": grado
        }
        
        # Also update the flat secciones for backward compatibility
        if "secciones" not in data:
            data["secciones"] = {}
        data["secciones"][seccion] = {
            "matriculados": sec_data.get("matriculados", sec_data["evaluados"]),
            "evaluados": sec_data["evaluados"],
            "AD_A": sec_data["AD_A"],
            "B": sec_data["B"],
            "C": sec_data["C"],
            "pct_AD_A": sec_data["pct_AD_A"],
            "pct_B": sec_data["pct_B"],
            "pct_C": sec_data["pct_C"]
        }
        
        total_evaluados += sec_data["evaluados"]
        total_sections += 1
        processed_sections.append(seccion)
        
        if sec_data["pct_C"] > 5.0:
            alerts.append(f"Seccion {seccion} tiene >{5}% de nivel C ({sec_data['pct_C']}%)")

    # Now rebuild grade totals from ALL sections uploaded for this bimester
    # (includes previously uploaded sections + newly uploaded ones)
    all_bim_sections = data["secciones_bimestres"].get(bimester_name, {})
    
    # Reset grade counts and rebuild from all available sections
    for g_key in grade_counts:
        grade_counts[g_key] = {"AD_A": 0, "B": 0, "C": 0, "evaluados": 0, "matriculados": 0}
    
    for sec_name, sec_info in all_bim_sections.items():
        grado = sec_info.get("grado", "")
        if grado in grade_counts:
            grade_counts[grado]["evaluados"] += sec_info["evaluados"]
            grade_counts[grado]["matriculados"] += sec_info.get("matriculados", sec_info["evaluados"])
            grade_counts[grado]["AD_A"] += sec_info["AD_A"]
            grade_counts[grado]["B"] += sec_info["B"]
            grade_counts[grado]["C"] += sec_info["C"]
        
    tot_eval = 0
    tot_ada = 0
    tot_b = 0
    tot_c = 0
    
    per_grade_summaries = {}
    
    for g, counts in grade_counts.items():
        e = counts["evaluados"]
        if e > 0:
            counts["pct_AD_A"] = round((counts["AD_A"] / e) * 100, 2)
            counts["pct_B"] = round((counts["B"] / e) * 100, 2)
            counts["pct_C"] = round((counts["C"] / e) * 100, 2)
        else:
            counts["pct_AD_A"] = 0.0
            counts["pct_B"] = 0.0
            counts["pct_C"] = 0.0
            
        tot_eval += e
        tot_ada += counts["AD_A"]
        tot_b += counts["B"]
        tot_c += counts["C"]
        
        if g not in data["grados"]:
            data["grados"][g] = {}
        data["grados"][g][bimester_name] = counts
        per_grade_summaries[g] = counts

    if tot_eval > 0:
        promedio_stats = {
            "evaluados": tot_eval,
            "AD_A": tot_ada,
            "B": tot_b,
            "C": tot_c,
            "pct_AD_A": round((tot_ada / tot_eval) * 100, 2),
            "pct_B": round((tot_b / tot_eval) * 100, 2),
            "pct_C": round((tot_c / tot_eval) * 100, 2)
        }
        if "Promedio General" not in data["grados"]:
            data["grados"]["Promedio General"] = {}
        data["grados"]["Promedio General"][bimester_name] = promedio_stats
        
    if bimester_name not in data["meta"]["periodos_disponibles"]:
        data["meta"]["periodos_disponibles"].append(bimester_name)
        tot_order = data["meta"].get("periodos_totales", ["Diagnóstica", "I Bimestre", "II Bimestre", "III Bimestre", "IV Bimestre"])
        data["meta"]["periodos_disponibles"].sort(key=lambda p: tot_order.index(p) if p in tot_order else 99)
        
    with open(json_path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
    
    all_uploaded = list(all_bim_sections.keys())
        
    detailed_results = {
        "total_evaluados": total_evaluados,
        "total_sections_now": total_sections,
        "total_sections_accumulated": len(all_uploaded),
        "all_uploaded_sections": all_uploaded,
        "newly_processed": processed_sections,
        "per_grade_summaries": per_grade_summaries,
        "per_section_summaries": sections_data,
        "alerts": alerts
    }
    
    return True, f"Bimestre '{bimester_name}': {total_sections} secciones procesadas ({total_evaluados} estudiantes). Total acumulado: {len(all_uploaded)}/10 secciones.", detailed_results


def get_upload_status(json_path="evaluaciones_2026.json"):
    """
    Returns the upload status for all bimesters showing which sections have been uploaded.
    """
    if not os.path.exists(json_path):
        return {}
    
    with open(json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    
    all_sections = ["1° A", "1° B", "2° A", "2° B", "3° A", "3° B", "4° A", "4° B", "5° A", "5° B"]
    bim_status = {}
    
    secciones_bim = data.get("secciones_bimestres", {})
    for bim_name, bim_secs in secciones_bim.items():
        uploaded = list(bim_secs.keys())
        pending = [s for s in all_sections if s not in uploaded]
        bim_status[bim_name] = {
            "uploaded": uploaded,
            "pending": pending,
            "count": len(uploaded),
            "total": len(all_sections),
            "complete": len(uploaded) == len(all_sections)
        }
    
    return bim_status


if __name__ == "__main__":
    success, msg = import_bimester_excel("calcular.xlsx", "II Bimestre")
    print(msg)
